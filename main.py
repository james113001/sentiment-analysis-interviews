import os
import json
import re
import polars as pl
from docx import Document
from ollama import chat
from nltk import sent_tokenize


# --------------------------------------------------------
# 1. Load transcript files (.txt or .docx)
# --------------------------------------------------------

def load_transcript(path: str) -> str:
    if path.endswith(".txt"):
        with open(path, "r", encoding="utf-8") as f:
            return f.read()

    elif path.endswith(".docx"):
        doc = Document(path)
        return "\n".join([p.text for p in doc.paragraphs])

    else:
        raise ValueError(f"Unsupported file format: {path}")


# --------------------------------------------------------
# 2. Remove interviewer lines
# --------------------------------------------------------

def extract_interviewee_text(full_text: str) -> str:
    cleaned_lines = []
    for line in full_text.split("\n"):
        line = line.strip()
        if not line:
            continue

        # Match participant/interviewee labels
        if re.match(r"(?i)^(participant|interviewee|p|resp|r)\s*[:\-]", line):
            cleaned_lines.append(
                re.sub(r"(?i)^(participant|interviewee|p|resp|r)\s*[:\-]\s*", "", line)
            )
        
        # Skip interviewer lines
        elif re.match(r"(?i)^(interviewer|int|i)\s*[:\-]", line):
            continue
        
        # If your transcripts do not label speakers, uncomment this:
        # cleaned_lines.append(line)

    return "\n".join(cleaned_lines)


# --------------------------------------------------------
# 3. Load themes sheet
# --------------------------------------------------------

def load_themes(theme_path: str) -> str:
    if theme_path.endswith(".xlsx"):
        df = pl.read_excel(theme_path)
    else:
        df = pl.read_csv(theme_path)

    lines = []
    for row in df.iter_rows(named=True):
        lines.append(f"- {row['code']}: {row['definition']}")

    return "\n".join(lines)


# --------------------------------------------------------
# 4. Code transcript using themes (Mistral)
# --------------------------------------------------------

def code_transcript_with_themes(text: str, theme_prompt: str):
    system_prompt = f"""
You are a qualitative analysis assistant. You will extract quotes from the transcript
and assign theme codes.

RULES:
- Cover 100% of interviewee content.
- Each sentence must appear in exactly one quote.
- A quote can contain 1–many sentences if they express the same theme.
- Do NOT omit or combine unrelated content.
- Do NOT include any interviewer lines.
- Return valid JSON only: list of {{"quote", "theme", "explanation"}}.

THEMES:
{theme_prompt}
"""

    user_prompt = f"Here is the full interviewee-only transcript:\n\n{text}\n\nExtract all theme-coded quotes now."

    response = chat(
        model="mistral",
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt}
        ]
    )

    cleaned = response["message"]["content"].strip()
    return json.loads(cleaned)


# --------------------------------------------------------
# 5. Process each transcript → 1 CSV per transcript
# --------------------------------------------------------

def process_folder(transcript_dir: str, theme_path: str, output_dir: str):
    os.makedirs(output_dir, exist_ok=True)

    theme_prompt = load_themes(theme_path)

    for filename in os.listdir(transcript_dir):
        if not (filename.endswith(".txt") or filename.endswith(".docx")):
            continue

        print(f"\n=== Processing {filename} ===")

        full_path = os.path.join(transcript_dir, filename)

        raw = load_transcript(full_path)
        participant_text = extract_interviewee_text(raw)

        coded_items = code_transcript_with_themes(participant_text, theme_prompt)

        # Convert to DataFrame
        df = pl.DataFrame(coded_items)

        out_name = f"coded_{os.path.splitext(filename)[0]}.csv"
        out_path = os.path.join(output_dir, out_name)

        df.write_csv(out_path)

        print(f"Saved → {out_path}")


# --------------------------------------------------------
# 6. Run
# --------------------------------------------------------

if __name__ == "__main__":
    process_folder(
        transcript_dir="transcripts",
        theme_path="themes.xlsx",
        output_dir="coded_output"
    )